#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Archive selected Outlook.com emails to PDF, attachments, CSV and Excel."""

import argparse
import base64
import csv
import hashlib
import importlib
import imaplib
import json
import os
import re
import shutil
import ssl
import sys
import textwrap
from datetime import datetime
from email import message_from_bytes, policy
from email.header import decode_header, make_header
from email.utils import getaddresses, parseaddr, parsedate_to_datetime
from getpass import getpass
from html import unescape as html_unescape
from xml.sax.saxutils import escape as xml_escape

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
except ImportError:
    arabic_reshaper = None
    get_display = None

def default_data_dir():
    if getattr(sys, 'frozen', False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


SCRIPT_DIR = default_data_dir()

CSV_COLUMNS = [
    'Ref No.', 'Agency', 'Subject', 'Transaction Date', 'Action', 'Response Date',
    'Sender Name', 'Sender Email', 'Recipients', 'CC', 'Attachments', 'Email PDF', 'Notes',
]
FILL_MANUALLY = '(fill manually)'
DASH = '\u2014'
NO_ATTACHMENT = '(no attachment)'
COLUMN_ALIASES = {
    'Ref No.': ('refno', 'ref', 'reference', 'refnumber', 'reference no', 'reference_number'),
    'Agency': ('agency', 'entity', 'authority'),
    'Subject': ('subject', 'topic', 'title'),
    'Transaction Date': ('transactiondate', 'date', 'transdate', 'transaction date'),
    'Action': ('action', 'status', 'state'),
    'Response Date': ('responsedate', 'replydate', 'response date', 'reply date'),
    'Sender Name': ('sendername', 'sender', 'fromname', 'from', 'from name'),
    'Sender Email': ('senderemail', 'fromemail', 'email', 'from email', 'mail'),
    'Recipients': ('recipients', 'to', 'toaddresses', 'to addresses'),
    'CC': ('cc', 'carboncopy'),
    'Attachments': ('attachments', 'attachment', 'files'),
    'Email PDF': ('emailpdf', 'pdf', 'pdfpath', 'pdf path', 'email pdf'),
    'Notes': ('notes', 'remarks', 'remark', 'comment', 'comments', 'note'),
}
ONEDRIVE_JUNK_RE = re.compile(
    r'1drv\.ms|res\.public\.onecdn\.static\.microsoft|res-1\.cdn\.office\.net'
    r'|mediap\.svc\.ms|microsoftpersonalcontent\.com|\.svc\.ms/transform'
    r'|fluentui-resources|tempauth', re.IGNORECASE)
ONEDRIVE_STRONG_RE = re.compile(
    r'res\.public\.onecdn\.static\.microsoft|res-1\.cdn\.office\.net'
    r'|mediap\.svc\.ms|microsoftpersonalcontent\.com|fluentui-resources'
    r'|cloud_blue_16x16|file-icon|^\s*\[', re.IGNORECASE)
ONEDRIVE_WEAK_RE = re.compile(r'1drv\.ms|tempauth', re.IGNORECASE)
BODY_WRAP_WIDTH = 70
VALUE_WRAP_WIDTH = 60
HEADER_FETCH_QUERY = '(UID INTERNALDATE FLAGS BODY.PEEK[HEADER.FIELDS (FROM TO CC SUBJECT DATE MESSAGE-ID REFERENCES IN-REPLY-TO)] BODYSTRUCTURE)'
FULL_FETCH_QUERY = '(RFC822)'
OAUTH_SCOPES = ['https://outlook.office.com/IMAP.AccessAsUser.All']

DEFAULT_CONFIG = {
    'email': 'your_email@outlook.com',
    'client_id': '',
    'tenant': 'common',
    'app_password': 'your_app_password_here',
    'imap_server': 'outlook.office365.com',
    'imap_port': 993,
    'use_ssl': True,
    'imap_provider': 'outlook',
    'fetch_limit': 50,
    'thread_view': True,
    'dark_mode': False,
    'exclude_keywords': [],
    'archive_folder': SCRIPT_DIR,
    'save_eml': True,
    'agency_map': {
        'dta.gov.sa': '\u0648\u0643\u0627\u0644\u0629 \u0627\u0644\u062a\u062d\u0648\u0644 \u0627\u0644\u0631\u0642\u0645\u064a \u0648\u0627\u0644\u0645\u062f\u0646 \u0627\u0644\u0630\u0643\u064a\u0629',
        'other.gov.sa': 'Other Government Agency',
    },
    'default_priority': 'Medium',
}

ARABIC_RE = re.compile(
    '[\u0600-\u06ff\u0750-\u077f\u08a0-\u08ff\ufb50-\ufdff\ufe70-\ufeff]'
)
SEQ_RE = re.compile(rb'^\s*(\d+)\s+FETCH\b')
UID_RE = re.compile(rb'\bUID\s+(\d+)')
INTERNALDATE_RE = re.compile(rb'INTERNALDATE\s+"([^"]+)"')
MESSAGE_ID_RE = re.compile(r'<([^<>]+)>')
REPLY_PREFIX_RE = re.compile(r'^(?:(?:re|fwd?|aw|wg|antw|sv|rv)\s*[:：]\s*|رد\s*[:：]?\s*)+', re.I)
BRACKET_TAG_RE = re.compile(r'^\s*(?:\[[^\]]*\]\s*)+')
STRUCT_NAME_RE = re.compile(r'"?NAME"?\s+"([^"]+)"', re.I)
STRUCT_ATTACHMENT_RE = re.compile(r'"ATTACHMENT"', re.I)
SEEN_RE = re.compile(rb'\\Seen', re.I)
HAS_FLAGS_RE = re.compile(rb'\bFLAGS\b')
REF_RE = re.compile(r'^GOV-(\d{4})-(\d+)$')
INTERNALDATE_FORMAT_RE = re.compile(r'(\d{1,2})-([A-Za-z]{3})-(\d{4})')
MONTHS = {
    'Jan': 1, 'Feb': 2, 'Mar': 3, 'Apr': 4, 'May': 5, 'Jun': 6,
    'Jul': 7, 'Aug': 8, 'Sep': 9, 'Oct': 10, 'Nov': 11, 'Dec': 12,
}
WINDOWS_RESERVED = {
    'con', 'prn', 'aux', 'nul',
    'com1', 'com2', 'com3', 'com4', 'com5', 'com6', 'com7', 'com8', 'com9',
    'lpt1', 'lpt2', 'lpt3', 'lpt4', 'lpt5', 'lpt6', 'lpt7', 'lpt8', 'lpt9',
}
COMMON_EXTENSIONS = {
    'octet-stream': 'bin', 'plain': 'txt', 'html': 'htm', 'jpeg': 'jpg',
    'jpg': 'jpg', 'png': 'png', 'gif': 'gif', 'pdf': 'pdf', 'csv': 'csv',
    'json': 'json', 'xml': 'xml', 'zip': 'zip', 'x-zip-compressed': 'zip',
    'x-msdownload': 'exe', 'vnd.openxml-officedocument.spreadsheetml.sheet': 'xlsx',
    'vnd.openxml-officedocument.wordprocessingml.document': 'docx',
    'vnd.ms-excel': 'xls', 'vnd.ms-word': 'doc', 'x-pkcs7-signature': 'p7s',
}
MAX_THREAD_FETCH = 15

_font_cache = {}


class ArchiveError(Exception):
    pass


def setup_console():
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8', errors='replace')
        except Exception:
            pass


def say(message=''):
    print(message)


def warn(message):
    print(f'[warning] {message}')


def fail(message):
    print(f'[error] {message}')


def ask(prompt):
    try:
        return input(prompt)
    except EOFError:
        raise ArchiveError('Input stream closed before the run finished.')


def _as_int(value, default):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _as_bool(value, default):
    if isinstance(value, bool):
        return value
    if isinstance(value, str):
        return value.strip().lower() in ('1', 'true', 'yes', 'y', 'on')
    if value is None:
        return default
    return bool(value)


def default_config_path():
    return os.path.join(SCRIPT_DIR, 'config.json')


def save_config(path, data):
    temporary = path + '.tmp'
    with open(temporary, 'w', encoding='utf-8') as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
        handle.write('\n')
    os.replace(temporary, path)


def load_config(path):
    if not os.path.exists(path):
        parent = os.path.dirname(path)
        try:
            if parent:
                os.makedirs(parent, exist_ok=True)
            save_config(path, DEFAULT_CONFIG)
        except OSError as exc:
            raise ArchiveError(f'Cannot create {path}: {exc}')
        say(f'Created a config file at {path}')
    try:
        with open(path, encoding='utf-8') as handle:
            data = json.load(handle)
    except json.JSONDecodeError as exc:
        raise ArchiveError(f'{path} is not valid JSON: {exc}')
    except OSError as exc:
        raise ArchiveError(f'Cannot read {path}: {exc}')
    if not isinstance(data, dict):
        raise ArchiveError(f'{path} must contain a JSON object, for example {{"email": "..."}}')
    data.setdefault('imap_server', 'outlook.office365.com')
    data['imap_port'] = _as_int(data.get('imap_port'), 993)
    data['use_ssl'] = _as_bool(data.get('use_ssl'), True)
    data['client_id'] = str(data.get('client_id') or '').strip()
    tenant = str(data.get('tenant') or 'common').strip() or 'common'
    if not re.fullmatch(r'[A-Za-z0-9-]+', tenant):
        raise ArchiveError('"tenant" in config.json must be "common", "consumers" or "organizations"')
    data['tenant'] = tenant
    data['fetch_limit'] = max(1, _as_int(data.get('fetch_limit'), 50))
    data['thread_view'] = _as_bool(data.get('thread_view'), True)
    data['dark_mode'] = _as_bool(data.get('dark_mode'), False)
    agency_map = data.get('agency_map')
    if not isinstance(agency_map, dict):
        agency_map = {}
    data['agency_map'] = {str(key).strip().lower(): str(value) for key, value in agency_map.items()}
    data['save_eml'] = _as_bool(data.get('save_eml'), True)
    exclude_keywords = data.get('exclude_keywords')
    if not isinstance(exclude_keywords, list):
        exclude_keywords = list(DEFAULT_CONFIG['exclude_keywords'])
    data['exclude_keywords'] = [str(word).strip() for word in exclude_keywords if str(word).strip()]
    data.setdefault('default_priority', 'Medium')
    return data


def keyword_matches(keyword, text):
    keyword = str(keyword).strip()
    if not keyword:
        return False
    if keyword.isascii():
        pattern = r'(?<![A-Za-z0-9])' + re.escape(keyword) + r'(?![A-Za-z0-9])'
        return re.search(pattern, text, re.IGNORECASE) is not None
    return keyword.casefold() in text.casefold()


def filter_excluded(summaries, config):
    keywords = config.get('exclude_keywords') or []
    if not keywords:
        return summaries, []
    kept = []
    dropped = []
    for item in summaries:
        haystack = ' '.join(
            str(part or '') for part in
            (item.get('from_name'), item.get('from_email'), item.get('subject'))
        )
        if any(keyword_matches(word, haystack) for word in keywords):
            dropped.append(item)
        else:
            kept.append(item)
    return kept, dropped


def ensure_credentials(config, config_path):
    placeholders_email = ('', 'your_email@outlook.com')
    placeholders_password = ('', 'your_app_password_here')
    changed = False
    save_password = False
    email_address = str(config.get('email') or '').strip()
    if email_address in placeholders_email:
        email_address = ask('Email address: ').strip()
        if not email_address:
            raise ArchiveError('No email address was entered.')
        config['email'] = email_address
        changed = True
    password = str(config.get('app_password') or '')
    if password in placeholders_password and not config.get('client_id'):
        try:
            password = getpass('App password (Microsoft account > Security > App passwords): ')
        except Exception:
            password = ask('App password: ')
        if not password:
            raise ArchiveError('No app password was entered.')
        config['app_password'] = password
        changed = True
        answer = ask('Save the app password in config.json for future runs? [y/N]: ').strip().lower()
        save_password = answer in ('y', 'yes')
        if not save_password:
            warn('The app password will not be saved - you will be asked again next run.')
    if changed:
        to_write = dict(config)
        if not save_password:
            to_write['app_password'] = 'your_app_password_here'
        try:
            save_config(config_path, to_write)
        except OSError as exc:
            warn(f'Could not update {config_path}: {exc}')


def resolve_root(config):
    raw = str(config.get('archive_folder') or '.').strip() or '.'
    if not os.path.isabs(raw):
        raw = os.path.join(SCRIPT_DIR, raw)
    return os.path.normpath(raw)


def read_config_quiet():
    path = default_config_path()
    try:
        with open(path, encoding='utf-8') as handle:
            return json.load(handle)
    except (OSError, json.JSONDecodeError):
        return {}


def read_log_rows(root):
    path = os.path.join(root, 'log.csv')
    if not os.path.exists(path):
        return []
    try:
        with open(path, newline='', encoding='utf-8-sig') as handle:
            return [dict(row) for row in csv.DictReader(handle)]
    except OSError:
        return []


def log_row_sort_key(row, column):
    value = str(row.get(column) or '').strip()
    if column in ('Ref No.', 'Transaction Date', 'Response Date'):
        numbers = re.findall(r'\d+', value)
        if numbers:
            return (0, tuple(int(part) for part in numbers), value.casefold())
    return (1, value.casefold(), value.casefold())


def filter_log_rows(rows, text='', date_from='', date_to='', agency='',
                    action='', attachment=None):
    text = str(text or '').strip().casefold()
    date_from = str(date_from or '').strip()
    date_to = str(date_to or '').strip()
    agency = str(agency or '').strip()
    action = str(action or '').strip()

    def keep(row):
        if text and text not in ' '.join(str(value or '') for value in row.values()).casefold():
            return False
        date_value = str(row.get('Transaction Date') or '').strip()
        if date_from and date_value and date_value < date_from:
            return False
        if date_to and date_value and date_value > date_to:
            return False
        if agency and str(row.get('Agency') or '').strip() != agency:
            return False
        if action and str(row.get('Action') or '').strip() != action:
            return False
        if attachment is not None:
            names = str(row.get('Attachments') or '').strip()
            has_attachment = bool(names) and names != DASH
            if has_attachment != attachment:
                return False
        return True

    return [row for row in rows if keep(row)]


def archive_stats(rows):
    total = len(rows)
    refs = set()
    agencies = {}
    months = {}
    with_attachments = 0
    outstanding = []
    actioned = []
    for row in rows:
        ref = str(row.get('Ref No.') or '').strip()
        if ref and ref != DASH:
            refs.add(ref)
        agency = str(row.get('Agency') or '').strip() or DASH
        agencies[agency] = agencies.get(agency, 0) + 1
        date_value = str(row.get('Transaction Date') or '').strip()
        month = date_value[:7] if len(date_value) >= 7 and date_value[4] == '-' else DASH
        months[month] = months.get(month, 0) + 1
        names = str(row.get('Attachments') or '').strip()
        if names and names != DASH:
            with_attachments += 1
        action = str(row.get('Action') or '').strip()
        response = str(row.get('Response Date') or '').strip()
        placeholders = (DASH, FILL_MANUALLY)
        if not action or action in placeholders or not response or response in placeholders:
            outstanding.append(row)
        else:
            actioned.append(row)
    return {
        'total': total,
        'threads': len(refs),
        'agencies': agencies,
        'months': months,
        'with_attachments': with_attachments,
        'outstanding': outstanding,
        'actioned': len(actioned),
    }


def remove_local_data(delete_archives):
    config = read_config_quiet()
    root = resolve_root(config)
    removed = []
    for name in ('config.json', 'token_cache.json', 'token_cache.json.bak'):
        path = os.path.join(SCRIPT_DIR, name)
        if os.path.isfile(path):
            try:
                os.remove(path)
                removed.append(path)
            except OSError:
                pass
    if delete_archives:
        for name in ('log.csv', 'log.xlsx', 'seen_ids.json'):
            path = os.path.join(root, name)
            if os.path.isfile(path):
                try:
                    os.remove(path)
                    removed.append(path)
                except OSError:
                    pass
        for name in ('PDF', 'EML', 'Attachments'):
            path = os.path.join(root, name)
            if os.path.isdir(path):
                try:
                    shutil.rmtree(path)
                    removed.append(path)
                except OSError:
                    pass
    return removed


def schedule_self_removal():
    if not getattr(sys, 'frozen', False):
        return False
    executable = os.path.abspath(sys.executable)
    if not executable.lower().endswith('.exe'):
        return False
    folder = os.path.dirname(executable)
    folders = [folder]
    parts = [f'del /f /q "{executable}"']
    launcher = os.environ.get('EAT_LAUNCHER')
    if launcher and launcher.lower().endswith('.exe'):
        launcher = os.path.abspath(launcher)
        if launcher != executable:
            parts.append(f'del /f /q "{launcher}"')
            parts.append(f'del /f /q "{os.path.join(os.path.dirname(launcher), "EmailArchiveTool_Manual.pdf")}"')
            folders.append(os.path.dirname(launcher))
    command = 'ping -n 3 127.0.0.1 >nul & ' + ' & '.join(parts + [f'rmdir "{d}"' for d in folders])
    try:
        import subprocess
        subprocess.Popen(command, shell=True, creationflags=0x00000008 | 0x08000000,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return True
    except Exception:
        return False


def ensure_dirs(root):
    try:
        for name in ('', 'PDF', 'Attachments', 'EML'):
            os.makedirs(os.path.join(root, name), exist_ok=True)
    except OSError as exc:
        raise ArchiveError(
            f'Cannot create the archive folder "{root}": {exc}\n'
            'Change "archive_folder" in config.json to a folder you can write to, '
            'for example "D:/GovEmailArchive".'
        )


def seen_path(root):
    return os.path.join(root, 'seen_ids.json')


def load_seen(root):
    path = seen_path(root)
    if not os.path.exists(path):
        return {}
    try:
        with open(path, encoding='utf-8') as handle:
            data = json.load(handle)
    except (OSError, json.JSONDecodeError):
        warn(f'Could not read {path} - already-archived emails may not be detected.')
        return {}
    if not isinstance(data, dict):
        return {}
    return {str(key): str(value) for key, value in data.items()}


def save_seen(root, seen):
    path = seen_path(root)
    temporary = path + '.tmp'
    try:
        with open(temporary, 'w', encoding='utf-8') as handle:
            json.dump(seen, handle, ensure_ascii=False, indent=2)
            handle.write('\n')
        os.replace(temporary, path)
    except OSError as exc:
        warn(f'Could not update {path}: {exc}')


def next_ref_no(root):
    year = datetime.now().year
    top = 0
    path = os.path.join(root, 'log.csv')
    if os.path.exists(path):
        try:
            with open(path, newline='', encoding='utf-8-sig') as handle:
                for row in csv.reader(handle):
                    if not row:
                        continue
                    match = REF_RE.match(str(row[0]).strip())
                    if match and int(match.group(1)) == year:
                        top = max(top, int(match.group(2)))
        except OSError:
            pass
    return f'GOV-{year}-{top + 1:03d}'


def append_log_row(root, row):
    path = os.path.join(root, 'log.csv')
    fresh = (not os.path.exists(path)) or os.path.getsize(path) == 0
    encoding = 'utf-8-sig' if fresh else 'utf-8'
    with open(path, 'a', newline='', encoding=encoding) as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, extrasaction='ignore')
        if fresh:
            writer.writeheader()
        writer.writerow({column: row.get(column, '') for column in CSV_COLUMNS})


def export_excel(root):
    csv_path = os.path.join(root, 'log.csv')
    if not os.path.exists(csv_path):
        raise ArchiveError('log.csv was not found - archive an email first.')
    try:
        with open(csv_path, newline='', encoding='utf-8-sig') as handle:
            rows = [list(row) for row in csv.reader(handle)]
    except OSError as exc:
        raise ArchiveError(f'Cannot read log.csv: {exc}')
    if not rows:
        raise ArchiveError('log.csv is empty - nothing to export.')
    width = len(CSV_COLUMNS)
    normalised = []
    for row in rows:
        trimmed = row[:width]
        trimmed += [''] * (width - len(trimmed))
        normalised.append(trimmed)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ArchiveError(f'openpyxl is not installed ({exc}) - run: python -m pip install -r requirements.txt')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Archive Log'
    for row in normalised:
        sheet.append(row)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column_index in range(1, sheet.max_column + 1):
        longest = len(str(normalised[0][column_index - 1] or ''))
        for row in normalised[1:]:
            value = str(row[column_index - 1] or '')
            visual = sum(2 if ord(char) > 0x0600 else 1 for char in value)
            longest = max(longest, min(visual, 60))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 10), 62)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    xlsx_path = os.path.join(root, 'log.xlsx')
    try:
        workbook.save(xlsx_path)
    except OSError as exc:
        raise ArchiveError(f'Cannot write log.xlsx: {exc}')
    say(f'Excel log exported: {xlsx_path} ({len(normalised) - 1} row(s))')
    return xlsx_path


def write_log_rows(root, rows):
    path = os.path.join(root, 'log.csv')
    tmp_path = path + '.tmp'
    with open(tmp_path, 'w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, extrasaction='ignore')
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, '') for column in CSV_COLUMNS})
    os.replace(tmp_path, path)
    xlsx_path = ''
    if rows:
        try:
            xlsx_path = export_excel(root)
        except ArchiveError:
            xlsx_path = ''
    return path, xlsx_path


def _norm_column(name):
    return re.sub(r'[^a-z0-9]', '', str(name or '').lower())


def import_column_map(fieldnames):
    available = {}
    for name in fieldnames or []:
        available.setdefault(_norm_column(name), name)
    mapping = {}
    for column in CSV_COLUMNS:
        candidates = (_norm_column(column),) + tuple(
            _norm_column(alias) for alias in COLUMN_ALIASES.get(column, ()))
        for candidate in candidates:
            if candidate in available:
                mapping[column] = available[candidate]
                break
    return mapping


def parse_import_csv(path):
    for encoding in ('utf-8-sig', 'cp1252', 'utf-8'):
        try:
            with open(path, newline='', encoding=encoding) as handle:
                reader = csv.DictReader(handle)
                mapping = import_column_map(reader.fieldnames)
                rows = []
                for raw in reader:
                    row = {}
                    for column in CSV_COLUMNS:
                        source = mapping.get(column)
                        value = str(raw.get(source) or '').strip() if source else ''
                        row[column] = value or DASH
                    rows.append(row)
                return rows
        except UnicodeDecodeError:
            continue
    raise ArchiveError(f'Cannot decode CSV: {path}')


def merge_log_rows(existing, incoming):
    def key_of(row):
        ref = str(row.get('Ref No.') or '').strip()
        return ref.casefold() if ref and ref != DASH else ''

    seen = {key_of(row) for row in existing}
    seen.discard('')
    merged = list(existing)
    added = 0
    skipped = 0
    for row in incoming:
        key = key_of(row)
        if key and key in seen:
            skipped += 1
            continue
        if key:
            seen.add(key)
        merged.append(row)
        added += 1
    return merged, added, skipped


def _styled_workbook(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Archive Log'
    sheet.append(CSV_COLUMNS)
    for row in rows:
        sheet.append([row.get(column, '') for column in CSV_COLUMNS])
    header_fill = PatternFill('solid', fgColor='1F4E78')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column_index in range(1, len(CSV_COLUMNS) + 1):
        column = CSV_COLUMNS[column_index - 1]
        longest = len(column)
        for row in rows:
            value = str(row.get(column) or '')
            visual = sum(2 if ord(char) > 0x0600 else 1 for char in value)
            longest = max(longest, min(visual, 60))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(longest + 2, 10), 62)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    return workbook


def export_rows_csv(path, rows):
    with open(path, 'w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_COLUMNS, extrasaction='ignore')
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column, '') for column in CSV_COLUMNS})
    return path


def export_rows_excel(path, rows):
    try:
        workbook = _styled_workbook(rows)
    except ImportError as exc:
        raise ArchiveError(f'openpyxl is not installed ({exc}) - run: python -m pip install -r requirements.txt')
    try:
        workbook.save(path)
    except OSError as exc:
        raise ArchiveError(f'Cannot write {path}: {exc}')
    return path


def check_dependencies():
    required = [('bs4', 'beautifulsoup4'), ('reportlab', 'reportlab'), ('openpyxl', 'openpyxl')]
    missing = []
    for module_name, package_name in required:
        try:
            importlib.import_module(module_name)
        except ImportError:
            missing.append(package_name)
    if missing:
        fail('Missing Python packages: ' + ', '.join(missing))
        say('Install them with:  python -m pip install -r requirements.txt')
        return False
    if arabic_reshaper is None or get_display is None:
        warn('arabic-reshaper / python-bidi not installed - Arabic text inside PDFs may not display correctly.')
    return True


def token_cache_path():
    return os.path.join(SCRIPT_DIR, 'token_cache.json')


def xoauth2_raw(email_address, token):
    return f'user={email_address}\x01auth=Bearer {token}\x01\x01'.encode('utf-8')


def _msal_detail(result):
    if not isinstance(result, dict):
        return str(result)
    detail = str(result.get('error_description') or result.get('error') or '').strip()
    if not detail:
        detail = str(result)
    return detail.replace('\r', '').strip()[:600]


def _consumers_required(detail):
    text = str(detail)
    return 'AADSTS9002346' in text or '/consumers endpoint' in text


def get_access_token(config):
    client_id = str(config.get('client_id') or '').strip()
    if not client_id:
        return None, None
    try:
        import msal
    except ImportError as exc:
        raise ArchiveError(f'msal is not installed ({exc}) - run: python -m pip install -r requirements.txt')
    email_address = str(config.get('email') or '').strip()
    cache = msal.SerializableTokenCache()
    cache_file = token_cache_path()
    if os.path.exists(cache_file):
        try:
            with open(cache_file, encoding='utf-8') as handle:
                cache.deserialize(handle.read())
        except (OSError, ValueError):
            pass
    tenants = [str(config.get('tenant') or 'common').strip() or 'common']
    if tenants[0].lower() != 'consumers':
        tenants.append('consumers')
    result = None
    for tenant in tenants:
        application = msal.PublicClientApplication(
            client_id,
            authority=f'https://login.microsoftonline.com/{tenant}',
            token_cache=cache,
        )
        result = None
        accounts = application.get_accounts(username=email_address or None)
        if accounts:
            silent = application.acquire_token_silent(OAUTH_SCOPES, account=accounts[0])
            if isinstance(silent, dict) and 'access_token' in silent:
                result = silent
        if not result:
            try:
                flow = application.initiate_device_flow(scopes=OAUTH_SCOPES)
            except Exception as exc:
                flow = {'error_description': str(exc)}
            if not isinstance(flow, dict) or 'user_code' not in flow:
                detail = _msal_detail(flow)
                if _consumers_required(detail) and tenant != tenants[-1]:
                    continue
                raise ArchiveError(
                    'Cannot start Microsoft sign-in: ' + detail + '\n'
                    'Check that "client_id" in config.json is the Application (client) ID of an Azure app '
                    'registered for personal Microsoft accounts, and that "Allow public client flows" is '
                    'enabled on its Authentication page (README, section "Sign in with Microsoft (OAuth2)").'
                )
            say(flow.get('message') or f'Open {flow.get("verification_uri")} and enter code {flow.get("user_code")}')
            sys.stdout.flush()
            try:
                result = application.acquire_token_by_device_flow(flow)
            except KeyboardInterrupt:
                raise ArchiveError('Microsoft sign-in was cancelled.')
            if not (isinstance(result, dict) and 'access_token' in result):
                if _consumers_required(_msal_detail(result)) and tenant != tenants[-1]:
                    continue
                break
        break
    if cache.has_state_changed:
        try:
            with open(cache_file, 'w', encoding='utf-8') as handle:
                handle.write(cache.serialize())
        except OSError as exc:
            warn(f'Could not save the Microsoft sign-in token cache: {exc}')
    if not result or 'access_token' not in result:
        raise ArchiveError(
            f'Microsoft sign-in failed for {email_address}: ' + _msal_detail(result) + '\n'
            'If the message mentions public client flows, enable "Allow public client flows" for the '
            'Azure app (README, section "Sign in with Microsoft (OAuth2)").'
        )
    return result['access_token'], get_token_username(result)


def get_token_username(result):
    if not isinstance(result, dict):
        return None
    claims = result.get('id_token_claims') or {}
    username = claims.get('preferred_username') or claims.get('upn')
    if not username:
        account = result.get('account')
        if isinstance(account, list):
            username = account[0] if account else ''
        elif account:
            username = account
    username = str(username or '').strip()
    return username or None


def login_failed_message(email_address, config, exc):
    text = str(exc)
    if 'basic authentication is disabled' in text.lower():
        if not str(config.get('client_id') or '').strip():
            return (
                f'IMAP login failed for {email_address}: Microsoft has disabled password sign-in '
                'for Outlook.com (basic authentication is disabled).\n'
                'Fix: register a free Azure app once and put its Application (client) ID in '
                '"client_id" in config.json - README, section "Sign in with Microsoft (OAuth2)".'
            )
        return (
            f'Microsoft sign-in failed for {email_address}: {text}\n'
            'Check "client_id" in config.json - README, section "Sign in with Microsoft (OAuth2)".'
        )
    if 'authenticated but not connected' in text.lower():
        return (
            f'IMAP login failed for {email_address}: {text}\n'
            'The Microsoft sign-in itself succeeded, but Outlook refused the IMAP session.\n'
            'The tested workaround: open Outlook settings (https://outlook.com > Settings > '
            'Mail > Forwarding and IMAP) and turn "Let devices and apps use IMAP" OFF, then '
            'ON again, wait about a minute, and retry fetching. No other change is needed - '
            'if it still fails, delete token_cache.json next to the program and try again, '
            'signing in with exactly this account.'
        )
    return (
        f'IMAP login failed for {email_address}: {text}\n'
        'Check that the app password is correct (a regular Microsoft password will not work), '
        'that IMAP is enabled in Outlook Settings > Mail > Sync email, '
        'and that the email address is correct.'
    )


def connect(config):
    server = str(config.get('imap_server') or 'outlook.office365.com')
    port = _as_int(config.get('imap_port'), 993)
    email_address = str(config.get('email') or '').strip()
    token, token_user = get_access_token(config)
    context = ssl.create_default_context()
    try:
        if config.get('use_ssl', True):
            connection = imaplib.IMAP4_SSL(server, port, ssl_context=context, timeout=60)
        else:
            connection = imaplib.IMAP4(server, port, timeout=60)
            try:
                connection.starttls(ssl_context=context)
            except (imaplib.IMAP4.error, OSError):
                pass
    except OSError as exc:
        raise ArchiveError(f'Cannot connect to {server} port {port}: {exc}')
    try:
        if token:
            auth_email = token_user or email_address
            if token_user and str(token_user).lower() != email_address.lower():
                warn(f'The signed-in Microsoft account ({token_user}) differs from "email" in '
                     f'config.json ({email_address}) - using the signed-in account for IMAP.')
            raw_token = xoauth2_raw(auth_email, token)
            connection.authenticate('XOAUTH2', lambda challenge: raw_token if not challenge else None)
        else:
            connection.login(email_address, str(config.get('app_password') or ''))
    except (imaplib.IMAP4.error, OSError) as exc:
        try:
            connection.logout()
        except Exception:
            pass
        raise ArchiveError(login_failed_message(email_address, config, exc))
    return connection


def close_conn(connection):
    try:
        connection.logout()
    except Exception:
        try:
            connection.shutdown()
        except Exception:
            pass


def select_inbox(connection):
    try:
        typ, _ = connection.select('INBOX', readonly=True)
    except imaplib.IMAP4.error as exc:
        raise ArchiveError(f'Cannot open the INBOX folder: {exc}')
    if typ != 'OK':
        raise ArchiveError('Cannot open the INBOX folder on the server.')


def search_all_uids(connection):
    typ, data = connection.uid('SEARCH', None, 'ALL')
    if typ != 'OK':
        raise ArchiveError('The server rejected the mailbox search command.')
    tokens = []
    for entry in data or []:
        if not entry:
            continue
        if isinstance(entry, bytes):
            tokens.extend(entry.split())
        else:
            tokens.extend(str(entry).split())
    uids = []
    for token in tokens:
        text = token.decode('ascii', 'ignore') if isinstance(token, bytes) else str(token)
        if text.isdigit():
            uids.append(text)
    return uids


def search_uids(connection, limit):
    return search_all_uids(connection)[-limit:]


def parse_internaldate(raw):
    text = raw.decode('ascii', 'ignore') if isinstance(raw, bytes) else str(raw)
    match = INTERNALDATE_FORMAT_RE.search(text)
    if not match:
        return ''
    month = MONTHS.get(match.group(2).capitalize(), 0)
    if not month:
        return ''
    return f'{match.group(3)}-{month:02d}-{int(match.group(1)):02d}'


def parse_fetch_records(data):
    records = []
    current = None
    for item in data or []:
        if item is None:
            continue
        if isinstance(item, tuple):
            current = {'meta': item[0] or b'', 'literal': item[1] or b'', 'tail': b''}
            records.append(current)
        elif SEQ_RE.match(item):
            current = {'meta': item, 'literal': b'', 'tail': b''}
            records.append(current)
        elif current is not None:
            current['tail'] += item
    for record in records:
        meta = record['meta']
        record['structure'] = meta + record['tail']
        uid_match = UID_RE.search(meta) or UID_RE.search(record['structure'])
        record['uid'] = uid_match.group(1).decode('ascii') if uid_match else ''
        date_match = INTERNALDATE_RE.search(meta)
        record['internal_date'] = parse_internaldate(date_match.group(1)) if date_match else ''
    return records


def meta_and_tail(record):
    return (record.get('meta') or b'') + (record.get('tail') or b'')


def fetch_summaries(connection, uids):
    typ, data = connection.uid('FETCH', ','.join(uids), HEADER_FETCH_QUERY)
    if typ != 'OK':
        raise ArchiveError('The server did not return the email headers.')
    requested = set(uids)
    summaries = []
    for index, record in enumerate(parse_fetch_records(data)):
        uid = record['uid']
        if uid not in requested:
            uid = uids[index] if index < len(uids) else ''
        if not uid:
            warn('Skipped an email whose message id could not be determined.')
            continue
        record['uid'] = uid
        if not record['literal']:
            warn('Skipped an email whose headers could not be read.')
            continue
        try:
            header = message_from_bytes(record['literal'], policy=policy.default)
        except Exception:
            warn('Skipped an email with an invalid header block.')
            continue
        from_name, from_email = split_address(str(header.get('from') or ''))
        subject = str(header.get('subject') or '').strip()
        flag_bytes = meta_and_tail(record)
        if HAS_FLAGS_RE.search(flag_bytes):
            seen = bool(SEEN_RE.search(flag_bytes))
        else:
            seen = True
        item = {
            'uid': record['uid'],
            'date': format_date(str(header.get('date') or '')) or record['internal_date'],
            'from_name': from_name,
            'from_email': from_email,
            'subject': subject,
            'message_id': str(header.get('message-id') or '').strip(),
            'in_reply_to': first_message_id(header.get('in-reply-to')),
            'references': message_id_list(header.get('references')),
            'attachment_names': [],
            'has_attachment': False,
            'seen': seen,
            'archived': False,
        }
        item['attachment_names'], item['has_attachment'] = attachment_names_from_structure(record['structure'])
        item['key'] = message_key(item)
        summaries.append(item)
    return summaries


def fetch_full(connection, uid):
    uid = str(uid or '').strip()
    if not uid.isdigit():
        raise ArchiveError('this email has no valid message id (UID) - fetch the list again and retry')
    typ, data = connection.uid('FETCH', uid, FULL_FETCH_QUERY)
    if typ != 'OK':
        raise ArchiveError(f'the server refused to download email UID {uid}')
    for record in parse_fetch_records(data):
        if record['literal']:
            return record['literal']
    raise ArchiveError(f'the server returned no content for email UID {uid}')


def attachment_names_from_structure(structure):
    if isinstance(structure, bytes):
        text = structure.decode('latin-1', 'replace')
    else:
        text = str(structure)
    names = []
    for match in STRUCT_NAME_RE.finditer(text):
        name = try_decode_words(match.group(1)).strip()
        if name and name not in names:
            names.append(name)
    has_attachment = bool(names) or bool(STRUCT_ATTACHMENT_RE.search(text))
    return names, has_attachment


def try_decode_words(text):
    if not text or '?=' not in text:
        return text
    try:
        return str(make_header(decode_header(text)))
    except Exception:
        return text


def split_address(raw):
    if not raw:
        return '', ''
    try:
        name, addr = parseaddr(raw)
    except Exception:
        return raw.strip(), ''
    return (name or '').strip(), (addr or '').strip()


def format_address_list(headers):
    parts = []
    for header in headers or []:
        try:
            entries = getaddresses([str(header)])
        except Exception:
            entries = []
        for name, addr in entries:
            name = (name or '').strip()
            addr = (addr or '').strip()
            if name:
                parts.append(name)
            elif addr:
                parts.append(addr)
    return '; '.join(parts) if parts else DASH


def format_date(raw):
    if not raw:
        return ''
    try:
        parsed = parsedate_to_datetime(raw)
        if parsed is not None:
            return parsed.strftime('%Y-%m-%d')
    except Exception:
        pass
    match = re.search(r'\d{4}-\d{2}-\d{2}', raw)
    return match.group(0) if match else ''


def message_key(item):
    message_id = (item.get('message_id') or '').strip()
    if message_id:
        return message_id
    basis = '|'.join([
        str(item.get('from_email') or ''),
        str(item.get('date') or ''),
        str(item.get('subject') or ''),
    ])
    return 'noid:' + hashlib.sha1(basis.encode('utf-8')).hexdigest()


def message_id_list(raw):
    if not raw:
        return []
    text = str(raw)
    found = [value.strip() for value in MESSAGE_ID_RE.findall(text) if value.strip()]
    if found:
        return found
    return [value for value in text.split() if value]


def normalize_message_id(value):
    text = str(value or '').strip()
    if text.startswith('<') and text.endswith('>') and len(text) > 2:
        text = text[1:-1].strip()
    return text


def first_message_id(raw):
    values = message_id_list(raw)
    return values[0] if values else ''


def normalize_subject(subject):
    text = BRACKET_TAG_RE.sub('', str(subject or '')).strip()
    while True:
        stripped = REPLY_PREFIX_RE.sub('', text).strip()
        if stripped == text:
            break
        text = stripped
    return re.sub(r'\s+', ' ', text).casefold().strip()


def is_reply_subject(subject):
    text = BRACKET_TAG_RE.sub('', str(subject or '')).strip()
    return bool(REPLY_PREFIX_RE.match(text))


def build_threads(items):
    count = len(items)
    parent = list(range(count))

    def find(node):
        while parent[node] != node:
            parent[node] = parent[parent[node]]
            node = parent[node]
        return node

    def union(left, right):
        root_left, root_right = find(left), find(right)
        if root_left != root_right:
            parent[root_right] = root_left

    id_index = {}
    for index, item in enumerate(items):
        message_id = normalize_message_id(item.get('message_id'))
        if message_id and message_id not in id_index:
            id_index[message_id] = index
    for index, item in enumerate(items):
        linked = [item.get('in_reply_to') or ''] + list(item.get('references') or [])
        for ref in linked:
            ref = normalize_message_id(ref)
            if ref and ref in id_index:
                union(index, id_index[ref])

    subject_groups = {}
    for index, item in enumerate(items):
        key = normalize_subject(item.get('subject'))
        if key:
            subject_groups.setdefault(key, []).append(index)
    for members in subject_groups.values():
        if len(members) < 2:
            continue
        if not any(is_reply_subject(items[i].get('subject')) for i in members):
            continue
        for other in members[1:]:
            union(members[0], other)

    groups = {}
    order = []
    for index in range(count):
        root = find(index)
        if root not in groups:
            groups[root] = []
            order.append(root)
        groups[root].append(index)
    thread_no = 0
    for root in order:
        members = groups[root]
        if len(members) > 1:
            thread_no += 1
            for position, index in enumerate(members):
                items[index]['thread_no'] = thread_no
                items[index]['thread_size'] = len(members)
                items[index]['thread_first'] = position == 0
        else:
            items[members[0]]['thread_no'] = 0
            items[members[0]]['thread_size'] = 1
            items[members[0]]['thread_first'] = False
    return [items[index] for root in order for index in groups[root]]


def thread_positions(summaries):
    mapping = {}
    for position, item in enumerate(summaries, 1):
        thread_no = item.get('thread_no') or 0
        if thread_no:
            mapping.setdefault(thread_no, []).append(position)
    return mapping


def missing_thread_refs(summaries):
    present = set()
    for item in summaries:
        message_id = normalize_message_id(item.get('message_id'))
        if message_id:
            present.add(message_id)
    wanted = []
    for item in summaries:
        candidates = []
        in_reply_to = normalize_message_id(item.get('in_reply_to'))
        if in_reply_to:
            candidates.append(in_reply_to)
        references = [normalize_message_id(value) for value in item.get('references') or []]
        references = [value for value in references if value]
        if references:
            candidates.append(references[0])
            if references[-1] != in_reply_to:
                candidates.append(references[-1])
        for ref in candidates:
            if ref and ref not in present and ref not in wanted:
                wanted.append(ref)
    return wanted[:MAX_THREAD_FETCH]


def fetch_missing_thread_messages(connection, summaries):
    refs = missing_thread_refs(summaries)
    if not refs:
        return []
    existing = {str(item.get('uid') or '') for item in summaries}
    found = []
    for ref in refs:
        try:
            typ, data = connection.uid('SEARCH', None, 'HEADER', 'Message-ID', f'"<{ref}>"')
        except Exception:
            continue
        if typ != 'OK' or not data:
            continue
        for entry in data:
            text = entry.decode('ascii', 'ignore') if isinstance(entry, bytes) else str(entry or '')
            for token in text.split():
                if token.isdigit() and token not in found and token not in existing:
                    found.append(token)
    if not found:
        return []
    try:
        return fetch_summaries(connection, found)
    except ArchiveError as exc:
        warn(f'Could not fetch older thread messages: {exc}')
        return []


def thread_sort_key(item):
    date = str(item.get('date') or '').strip() or '9999-99-99'
    uid = str(item.get('uid') or '').strip()
    return (date, int(uid) if uid.isdigit() else 0)


def agency_for(email_address, config):
    email_address = (email_address or '').strip().lower()
    if '@' not in email_address:
        return DASH
    domain = email_address.rsplit('@', 1)[1].strip('.')
    agency_map = config.get('agency_map') or {}
    if domain in agency_map:
        return agency_map[domain]
    for key, value in agency_map.items():
        if domain.endswith('.' + key):
            return value
    return DASH


def build_notes(from_name, from_email, subject):
    if from_name:
        who = from_name.split()[0]
    elif from_email:
        who = from_email.split('@')[0]
    else:
        who = 'unknown'
    if subject.lower().startswith('re:'):
        return f'Received from {who}, reply on same thread'
    if not subject or subject == DASH:
        return f'Received from {who}'
    return f'Received from {who}, subject: {subject}'


def decode_part(part):
    try:
        data = part.get_payload(decode=True)
    except Exception:
        data = None
    if data:
        charsets = [part.get_content_charset(), 'utf-8', 'cp1256', 'latin-1']
        for charset in charsets:
            if not charset:
                continue
            try:
                return data.decode(charset)
            except (LookupError, UnicodeDecodeError):
                continue
        return data.decode('utf-8', errors='replace')
    payload = part.get_payload()
    return payload if isinstance(payload, str) else ''


def split_parts(message):
    plain, html, attachments = [], [], []

    def walk(part):
        if part.get_content_maintype() == 'message':
            attachments.append(part)
            return
        if part.is_multipart():
            payload = part.get_payload()
            if isinstance(payload, list):
                for sub in payload:
                    walk(sub)
            return
        if part.get_content_disposition() == 'inline':
            return
        filename = part.get_filename()
        disposition = part.get_content_disposition()
        if filename or disposition == 'attachment':
            attachments.append(part)
            return
        content_type = part.get_content_type()
        if content_type == 'text/plain':
            plain.append(part)
        elif content_type == 'text/html':
            html.append(part)

    walk(message)
    return plain, html, attachments


def normalize_plain(text):
    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\xa0', ' ')
    lines = [line.rstrip() for line in text.split('\n')]
    output = []
    blanks = 0
    for line in lines:
        if line:
            output.append(line)
            blanks = 0
        else:
            blanks += 1
            if blanks == 1 and output:
                output.append('')
    return '\n'.join(output).strip()


def strip_onedrive_footer(text):
    """Remove Outlook web "insert link" clutter from a message body.

    Outlook.com pastes shared OneDrive files as a block of lines made of
    Microsoft icon/thumbnail/asset URLs, tempauth chains and 1drv.ms share
    links. A contiguous run of such lines is removed whenever it contains
    at least one strong (asset/icon/thumbnail) signature. A normal sentence
    that merely mentions a share link (weak signature only) is left intact.
    """
    lines = str(text or '').split('\n')
    if not any(ONEDRIVE_STRONG_RE.search(line) for line in lines):
        return normalize_plain(str(text)) if text != '' else ''
    output = []
    index = 0
    count = len(lines)
    while index < count:
        line = lines[index]
        is_junk = (line.strip()
                   and (ONEDRIVE_STRONG_RE.search(line) or ONEDRIVE_WEAK_RE.search(line)))
        if not is_junk:
            output.append(line)
            index += 1
            continue
        end = index
        strong = False
        while end < count and (not lines[end].strip()
                               or ONEDRIVE_STRONG_RE.search(lines[end])
                               or ONEDRIVE_WEAK_RE.search(lines[end])):
            if lines[end].strip() and ONEDRIVE_STRONG_RE.search(lines[end]):
                strong = True
            end += 1
        while end > index and not lines[end - 1].strip():
            end -= 1
        if strong:
            while end < count and not lines[end].strip():
                end += 1
            index = end
        else:
            output.extend(lines[index:end])
            index = end
    return normalize_plain('\n'.join(output))


def html_to_text(html_text):
    if not html_text:
        return ''
    if BeautifulSoup is None:
        return strip_html_tags(html_text)
    try:
        soup = BeautifulSoup(html_text, 'html.parser')
    except Exception:
        return strip_html_tags(html_text)
    for element in soup(['script', 'style', 'head', 'title']):
        element.decompose()
    for element in soup.find_all('br'):
        element.replace_with('\n')
    for element in soup.find_all(['p', 'div', 'tr', 'blockquote', 'section', 'article', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6']):
        element.append('\n')
    for element in soup.find_all('li'):
        element.insert(0, '- ')
        element.append('\n')
    for element in soup.find_all('a'):
        if ONEDRIVE_JUNK_RE.search(element.get('href') or ''):
            element.decompose()
    for element in soup.find_all('img'):
        if ONEDRIVE_JUNK_RE.search(element.get('src') or ''):
            element.decompose()
    return normalize_plain(soup.get_text())


def strip_html_tags(html_text):
    text = re.sub(r'(?is)<(script|style).*?>.*?</\1>', ' ', html_text)
    text = re.sub(r'(?i)<br\s*/?>', '\n', text)
    text = re.sub(r'(?i)</(p|div|tr|li|h[1-6])>', '\n', text)
    text = re.sub(r'<[^>]+>', ' ', text)
    return normalize_plain(html_unescape(text))


def extract_body(message):
    plain, html, _ = split_parts(message)
    chunks = [strip_onedrive_footer(normalize_plain(decode_part(part)))
              for part in plain]
    chunks = [chunk for chunk in chunks if chunk]
    if chunks:
        return '\n\n'.join(chunks)
    chunks = [strip_onedrive_footer(html_to_text(decode_part(part)))
              for part in html]
    chunks = [chunk for chunk in chunks if chunk]
    if chunks:
        return '\n\n'.join(chunks)
    return '(no message body)'


def eml_body(path):
    try:
        with open(path, 'rb') as handle:
            message = message_from_bytes(handle.read(), policy=policy.default)
    except (OSError, ValueError):
        return '(could not read the saved email)'
    try:
        return extract_body(message)
    except Exception as exc:
        return f'(could not read the message body: {exc})'


def build_record(message, config, fallback_date=''):
    from_name, from_email = split_address(str(message.get('from') or ''))
    subject = str(message.get('subject') or '').strip() or DASH
    date_text = format_date(str(message.get('date') or '')) or fallback_date
    if not date_text:
        date_text = datetime.now().strftime('%Y-%m-%d')
        warn(f'Email "{subject}" has no readable date - using today for Transaction Date.')
    recipients = format_address_list(message.get_all('to', []))
    cc = format_address_list(message.get_all('cc', []))
    return {
        'subject': subject,
        'date': date_text,
        'from_name': from_name,
        'from_email': from_email,
        'recipients': recipients,
        'cc': cc,
        'body': extract_body(message),
        'agency': agency_for(from_email, config),
        'notes': build_notes(from_name, from_email, subject),
    }


def build_row(ref, record, attachment_names, pdf_path):
    return {
        'Ref No.': ref,
        'Agency': record['agency'],
        'Subject': record['subject'],
        'Transaction Date': record['date'],
        'Action': FILL_MANUALLY,
        'Response Date': FILL_MANUALLY,
        'Sender Name': record['from_name'] or DASH,
        'Sender Email': record['from_email'] or DASH,
        'Recipients': record['recipients'],
        'CC': record['cc'],
        'Attachments': '; '.join(attachment_names) if attachment_names else DASH,
        'Email PDF': pdf_path,
        'Notes': record['notes'],
    }


def sanitize_filename(name):
    name = str(name or '')
    name = re.sub(r'[\x00-\x1f\x7f]', '', name)
    name = re.sub(r'[<>:"/\\|?*]', '_', name).strip(' .')
    if not name:
        name = 'attachment'
    stem, extension = os.path.splitext(name)
    if len(extension) > 16:
        stem = stem + extension[16:]
        extension = extension[:16]
    if len(stem) > 100:
        stem = stem[:100]
    if stem.lower() in WINDOWS_RESERVED:
        stem = '_' + stem
    return stem + extension


def default_attachment_name(part, index):
    if part.get_content_maintype() == 'message':
        return f'attachment_{index}.eml'
    subtype = (part.get_content_subtype() or '').lower()
    extension = COMMON_EXTENSIONS.get(subtype)
    if not extension:
        tail = subtype.split('.')[-1]
        extension = re.sub(r'[^a-z0-9]', '', tail)[:6] or 'bin'
    return f'attachment_{index}.{extension}'


def build_attachment_name(original, date_text):
    name = sanitize_filename(original)
    stem, extension = os.path.splitext(name)
    if date_text and not re.search(r'\d{4}-\d{2}-\d{2}', stem):
        stem = f'{stem}_{date_text}'
    return stem + extension


def name_candidates(preferred, ref, extension):
    stem, _ = os.path.splitext(preferred)
    yield preferred
    yield f'{stem}_{ref}{extension}'
    counter = 1
    while counter < 1000:
        yield f'{stem}_{ref}_{counter}{extension}'
        counter += 1


def attachment_bytes(part):
    try:
        data = part.get_payload(decode=True)
    except Exception:
        data = None
    if data:
        return data
    if part.get_content_maintype() == 'message':
        payload = part.get_payload()
        inner = payload[0] if isinstance(payload, list) and payload else None
        if inner is None:
            return None
        try:
            return inner.as_bytes(policy=policy.default)
        except Exception:
            return None
    try:
        payload = part.get_payload()
    except Exception:
        return None
    if isinstance(payload, str):
        return payload.encode('utf-8', 'replace')
    return None


def save_attachments(message, attachments_dir, ref, date_text):
    _, _, parts = split_parts(message)
    saved = []
    taken = set()
    for index, part in enumerate(parts, 1):
        original = part.get_filename() or default_attachment_name(part, index)
        preferred = build_attachment_name(original, date_text)
        _, extension = os.path.splitext(preferred)
        chosen = None
        for candidate in name_candidates(preferred, ref, extension):
            low = candidate.lower()
            if low not in taken and not os.path.exists(os.path.join(attachments_dir, candidate)):
                chosen = candidate
                break
        if chosen is None:
            warn(f'Skipped attachment (no free file name): {original}')
            continue
        data = attachment_bytes(part)
        if not data:
            warn(f'Skipped unsupported attachment: {original}')
            continue
        try:
            with open(os.path.join(attachments_dir, chosen), 'wb') as handle:
                handle.write(data)
        except OSError as exc:
            warn(f'Could not save attachment {original}: {exc}')
            continue
        taken.add(chosen.lower())
        saved.append(chosen)
    return saved


def parse_selection(text, count, threads=None):
    cleaned = text.strip().lower()
    if cleaned in ('q', 'quit', 'c', 'cancel', 'none'):
        return []
    if cleaned in ('all', 'a', '*'):
        return list(range(1, count + 1))
    numbers = []
    for token in re.split(r'[,\s;]+', cleaned):
        if not token:
            continue
        thread_match = re.fullmatch(r't(\d+)', token)
        if thread_match:
            if threads is None:
                raise ValueError('Thread selection is not enabled - "thread_view" in config.json')
            thread_no = int(thread_match.group(1))
            positions = threads.get(thread_no)
            if not positions:
                raise ValueError(f'Thread {thread_no} does not exist in the current list')
            for value in positions:
                if value not in numbers:
                    numbers.append(value)
            continue
        if not token.isdigit():
            raise ValueError(f'Invalid part "{token}" - use numbers like 1,3,5 or threads like t2')
        value = int(token)
        if value < 1 or value > count:
            raise ValueError(f'Number {value} is outside the list (1-{count})')
        if value not in numbers:
            numbers.append(value)
    if not numbers:
        raise ValueError('No numbers were entered.')
    return numbers


def wrap_text(text, width):
    lines = []
    normalised = str(text).replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')
    for line in normalised.split('\n'):
        stripped = line.rstrip()
        if not stripped:
            lines.append('')
            continue
        wrapped = textwrap.wrap(stripped, width=width, break_long_words=True, break_on_hyphens=False)
        lines.extend(wrapped if wrapped else [stripped[:width]])
    return lines


def shape_text(text):
    if arabic_reshaper is None or get_display is None:
        return text
    if not ARABIC_RE.search(text):
        return text
    try:
        return get_display(arabic_reshaper.reshape(text))
    except Exception:
        return text


def markup_line(text, width=BODY_WRAP_WIDTH):
    if text is None or text == '':
        return '&#160;'
    parts = []
    for line in wrap_text(text, width):
        if not line:
            parts.append('&#160;')
            continue
        parts.append(xml_escape(shape_text(line)))
    return '<br/>'.join(parts) if parts else '&#160;'


def register_fonts():
    if _font_cache:
        return _font_cache['names']
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except ImportError as exc:
        raise ArchiveError(f'reportlab is not installed ({exc}) - run: python -m pip install -r requirements.txt')
    windows_fonts = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts')
    candidates = [
        (os.path.join(windows_fonts, 'arial.ttf'), os.path.join(windows_fonts, 'arialbd.ttf')),
        (os.path.join(windows_fonts, 'tahoma.ttf'), os.path.join(windows_fonts, 'tahomabd.ttf')),
        (os.path.join(windows_fonts, 'segoeui.ttf'), os.path.join(windows_fonts, 'segoeuib.ttf')),
        ('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'),
        ('/usr/share/fonts/dejavu/DejaVuSans.ttf', '/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf'),
        ('/System/Library/Fonts/Supplemental/Arial.ttf', '/System/Library/Fonts/Supplemental/Arial Bold.ttf'),
    ]
    for regular, bold in candidates:
        if not os.path.exists(regular):
            continue
        bold_path = bold if os.path.exists(bold) else regular
        try:
            pdfmetrics.registerFont(TTFont('GovSans', regular))
            pdfmetrics.registerFont(TTFont('GovSans-Bold', bold_path))
            pdfmetrics.registerFontFamily('GovSans', normal='GovSans', bold='GovSans-Bold',
                                          italic='GovSans', boldItalic='GovSans-Bold')
            _font_cache['names'] = ('GovSans', 'GovSans-Bold')
            _font_cache['arabic_ok'] = True
            return _font_cache['names']
        except Exception:
            continue
    _font_cache['names'] = ('Helvetica', 'Helvetica-Bold')
    _font_cache['arabic_ok'] = False
    return _font_cache['names']


def make_pdf(path, ref, records):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise ArchiveError(f'reportlab is not installed ({exc}) - run: python -m pip install -r requirements.txt')
    font_name, bold_name = register_fonts()
    if not _font_cache.get('arabic_ok') and not _font_cache.get('arabic_warned'):
        if ARABIC_RE.search(records[0][0]['subject'] + records[0][0]['body']):
            _font_cache['arabic_warned'] = True
            warn('No Arabic-capable font was found on this computer - Arabic text in PDFs may not display correctly.')
    title_style = ParagraphStyle('GovTitle', fontName=bold_name, fontSize=13, leading=17,
                                 alignment=TA_CENTER, textColor=colors.HexColor('#12395e'), spaceAfter=2)
    subtitle_style = ParagraphStyle('GovSubtitle', fontName=font_name, fontSize=9.5, leading=13,
                                    alignment=TA_CENTER, textColor=colors.HexColor('#555555'), spaceAfter=8)
    ref_style = ParagraphStyle('GovRef', fontName=bold_name, fontSize=11, leading=15, spaceAfter=2)
    meta_style = ParagraphStyle('GovMeta', fontName=font_name, fontSize=10, leading=14, spaceAfter=2)
    heading_style = ParagraphStyle('GovHeading', fontName=bold_name, fontSize=11, leading=15,
                                   textColor=colors.HexColor('#12395e'), spaceBefore=10, spaceAfter=4)
    body_style = ParagraphStyle('GovBody', fontName=font_name, fontSize=10.5, leading=15, spaceAfter=4)
    divider_style = ParagraphStyle('GovDivider', fontName=bold_name, fontSize=10.5, leading=14,
                                   textColor=colors.HexColor('#12395e'), spaceAfter=5)
    story = [
        Paragraph(markup_line('GOVERNMENT EMAIL ARCHIVE RECORD'), title_style),
        Paragraph(markup_line('\u0633\u062c\u0644 \u0623\u0631\u0634\u0627\u0641\u0629 \u0627\u0644\u0628\u0631\u064a\u062f \u0627\u0644\u0625\u0644\u0643\u062a\u0631\u0648\u0646\u064a'), subtitle_style),
        Paragraph(f'<b>Ref No.:</b>&nbsp;{xml_escape(ref)}', ref_style),
        Spacer(1, 6),
    ]
    total = len(records)
    for index, (record, attachment_names) in enumerate(records):
        if total > 1:
            sender = record['from_name'] or record['from_email'] or DASH
            divider = f'Message {index + 1} of {total} \u2014 {record["date"]} \u2014 {sender}'
            story.append(Paragraph(markup_line(divider), divider_style))
        meta_fields = [
            ('Agency', record['agency']),
            ('Subject', record['subject']),
            ('Transaction Date', record['date']),
            ('Sender Name', record['from_name'] or DASH),
            ('Sender Email', record['from_email'] or DASH),
            ('Recipients', record['recipients']),
            ('CC', record['cc']),
            ('Attachments', '; '.join(attachment_names) if attachment_names else DASH),
        ]
        for label, value in meta_fields:
            story.append(Paragraph(f'<b>{label}:</b>&nbsp;{markup_line(value, VALUE_WRAP_WIDTH)}', meta_style))
        story.append(Paragraph('<b>Message body</b>', heading_style))
        story.append(Paragraph(markup_line(record['body']), body_style))
        if index < total - 1:
            story.append(PageBreak())

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor('#777777'))
        canvas.drawString(document.leftMargin, 1.1 * cm, 'Generated by Gov Email Archive Tool')
        canvas.drawRightString(A4[0] - document.rightMargin, 1.1 * cm, f'Page {canvas.getPageNumber()}')
        canvas.restoreState()

    document = SimpleDocTemplate(
        path, pagesize=A4,
        leftMargin=1.7 * cm, rightMargin=1.7 * cm, topMargin=1.7 * cm, bottomMargin=1.9 * cm,
        title=ref, author='Gov Email Archive Tool', subject='Archived email record',
    )
    try:
        document.build(story, onFirstPage=footer, onLaterPages=footer)
    except Exception as exc:
        raise ArchiveError(f'could not create the PDF for {ref}: {exc}')


def display_summaries(summaries):
    say('')
    say('  Recent emails (oldest first):')
    for index, item in enumerate(summaries, 1):
        if item.get('thread_first'):
            group = [member for member in summaries if member.get('thread_no') == item['thread_no']]
            done = sum(1 for member in group if member['archived'])
            done_text = f', {done} already archived' if done else ''
            thread_subject = item['subject'] or '(no subject)'
            if len(thread_subject) > 64:
                thread_subject = thread_subject[:63] + '...'
            say(f'  Thread {item["thread_no"]} ({len(group)} messages{done_text}): {thread_subject}')
        date_text = item['date'] or '----------'
        name = item['from_name'] or item['from_email'] or '(unknown sender)'
        if len(name) > 34:
            name = name[:33] + '...'
        subject = item['subject'] or '(no subject)'
        if len(subject) > 64:
            subject = subject[:63] + '...'
        names = item['attachment_names']
        if names:
            attachment_text = names[0] if len(names) == 1 else f'{len(names)} files: {names[0]} ...'
        elif item['has_attachment']:
            attachment_text = '(attachment)'
        else:
            attachment_text = NO_ATTACHMENT
        archived = '  [already archived]' if item['archived'] else ''
        say(f'[{index:>3}] {date_text} | {name} | {subject} | {attachment_text}{archived}')
    say('')


def keyword_for(item):
    name = str(item.get('from_name') or '').strip()
    if name:
        return name
    email = str(item.get('from_email') or '').strip()
    if email and '@' in email:
        return email
    return ''


def keywords_for(item):
    """Hide keywords representing a sender: its display name and, when
    different, its email address. Both go into exclude_keywords, so a hidden
    sender stays hidden even when its name is written differently on other
    messages, and an exact address never collides with a common name.
    """
    keywords = []
    name = str(item.get('from_name') or '').strip()
    if name:
        keywords.append(name)
    email = str(item.get('from_email') or '').strip()
    if email and '@' in email and email not in keywords:
        keywords.append(email)
    return keywords


def render_keyword_browser(items, checked, words):
    say('')
    say('=' * 66)
    say('  Exclude keywords - check emails to HIDE (saved as exclude_keywords in config.json)')
    say('  numbers toggle | "t2" whole thread | "k1" other keyword | "x" clear all')
    say('  Enter or "q" saves, nothing is ever discarded')
    say('=' * 66)
    for index, item in enumerate(items):
        marker = '[x]' if checked[index] else '[ ]'
        date_text = item['date'] or '----------'
        name = item['from_name'] or item['from_email'] or '(unknown sender)'
        if len(name) > 28:
            name = name[:27] + '...'
        subject = item['subject'] or '(no subject)'
        if len(subject) > 44:
            subject = subject[:43] + '...'
        thread = f'  [thread {item["thread_no"]}]' if item.get('thread_no') else ''
        hidden_keywords = keywords_for(item)
        hint = ('  -> hides ' + ', '.join('"%s"' % word for word in hidden_keywords)) \
            if hidden_keywords else '  -> no sender to filter on'
        say(f'  {marker} {index + 1:>3}) {date_text} | {name} | {subject}{thread}{hint}')
    extras = [word for word in words
              if not any(word == keyword for item in items for keyword in keywords_for(item))]
    if extras:
        say('')
        say('  Other saved keywords (not in the list above): '
            + '  '.join(f'k{i}: {word}' for i, word in enumerate(extras, 1)))
    say('')


def run_keyword_browser(items, words, prompt):
    keywords = [keywords_for(item) for item in items]
    checked = [any(keyword in words for keyword in item_keywords) for item_keywords in keywords]
    while True:
        render_keyword_browser(items, checked, words)
        text = (prompt('Toggle numbers (e.g. 1,3), "t2" thread, "k1" keyword, "x" clear, Enter or q saves: ') or '').strip()
        if not text:
            return words
        for token in re.split(r'[\s,;]+', text):
            token = token.strip()
            lowered = token.lower()
            if lowered in ('q', 'quit', 'cancel', 'save', 'done'):
                return words
            if lowered in ('a', 'all'):
                for index in range(len(items)):
                    checked[index] = not checked[index]
                    keywords_list = keywords[index]
                    if checked[index]:
                        for keyword in keywords_list:
                            if keyword and keyword not in words:
                                words.append(keyword)
                    else:
                        words[:] = [word for word in words if word not in keywords_list]
                continue
            if lowered in ('x', 'clear'):
                words[:] = []
                for index in range(len(items)):
                    checked[index] = False
                continue
            if lowered[0] == 'k' and lowered[1:].isdigit():
                extras = [word for word in words
                          if not any(word == keyword for keyword_list in keywords
                                     for keyword in keyword_list)]
                index = int(lowered[1:]) - 1
                if 0 <= index < len(extras):
                    target = extras[index]
                    if target in words:
                        words[:] = [word for word in words if word != target]
                    else:
                        words.append(target)
                else:
                    warn(f'Unknown keyword number {token}.')
                continue
            if lowered[0] == 't' and lowered[1:].isdigit():
                thread_no = int(lowered[1:])
                affected = [i for i, item in enumerate(items) if item.get('thread_no') == thread_no]
                if not affected:
                    warn(f'Unknown thread number {thread_no}.')
                    continue
                for index in affected:
                    checked[index] = not checked[index]
                    keywords_list = keywords[index]
                    if checked[index]:
                        for keyword in keywords_list:
                            if keyword and keyword not in words:
                                words.append(keyword)
                    else:
                        words[:] = [word for word in words if word not in keywords_list]
                continue
            if token.isdigit():
                index = int(token) - 1
                if not (0 <= index < len(items)):
                    warn(f'Unknown row number {token}.')
                    continue
                if not keywords[index]:
                    warn(f'Row {token} has no sender to filter on.')
                    continue
                checked[index] = not checked[index]
                keywords_list = keywords[index]
                if checked[index]:
                    for keyword in keywords_list:
                        if keyword and keyword not in words:
                            words.append(keyword)
                else:
                    words[:] = [word for word in words if word not in keywords_list]
                continue
            warn(f'Unknown command "{token}".')


def cmd_keywords(config_path, config, root, summaries):
    if not summaries:
        say('Fetching the email list first...')
        summaries = cmd_fetch(config, root)
        if not summaries:
            return
    words = list(config.get('exclude_keywords') or [])
    try:
        result = run_keyword_browser(summaries, words, ask)
    except KeyboardInterrupt:
        say('')
        warn('Cancelled - no changes were saved.')
        return
    if result is None:
        say('Cancelled - no changes were saved.')
        return
    config['exclude_keywords'] = result
    try:
        save_config(config_path, config)
        say(f'Saved {len(result)} keyword(s) to {os.path.basename(config_path)}.')
    except OSError as exc:
        warn(f'Could not save {config_path}: {exc}')


def cmd_fetch(config, root):
    server = str(config.get('imap_server'))
    say(f'Connecting to {server}...')
    connection = connect(config)
    try:
        select_inbox(connection)
        limit = config['fetch_limit']
        all_uids = search_all_uids(connection)
        if not all_uids:
            warn('The mailbox is empty - no emails to show.')
            return []
        summaries = []
        excluded_total = 0
        fetched = set()
        position = max(0, len(all_uids) - limit)
        while len(summaries) < limit and position >= 0:
            chunk = [uid for uid in all_uids[position:] if uid not in fetched]
            if not chunk:
                break
            fetched.update(chunk)
            say(f'Fetching details for {len(chunk)} email(s)...')
            batch = fetch_summaries(connection, chunk)
            batch, excluded = filter_excluded(batch, config)
            excluded_total += len(excluded)
            summaries = batch + summaries
            position -= limit
        if config.get('thread_view', True):
            extras = fetch_missing_thread_messages(connection, summaries)
            if extras:
                summaries = sorted(summaries + extras, key=thread_sort_key)
                say(f'Added {len(extras)} older message(s) to complete threads.')
                summaries, excluded = filter_excluded(summaries, config)
                excluded_total += len(excluded)
        if excluded_total:
            say(f'Skipped {excluded_total} excluded message(s) matching exclude_keywords in config.json.')
    finally:
        close_conn(connection)
    if not summaries:
        if excluded_total:
            warn('Every fetched message was excluded by exclude_keywords - adjust the list in config.json.')
        else:
            warn('No readable email headers were returned.')
        return []
    seen = load_seen(root)
    for item in summaries:
        item['archived'] = item['key'] in seen
    thread_view = bool(config.get('thread_view', True))
    if thread_view:
        summaries = build_threads(summaries)
    display_summaries(summaries)
    say(f'{len(summaries)} email(s) shown (fetch_limit={limit}).')
    if thread_view and any(item.get('thread_no') for item in summaries):
        say('Type the numbers to archive, e.g. 1,3,5  - or t2 to archive a whole thread  (or run choice 2).')
    else:
        say('Type the numbers to archive, e.g. 1,3,5  (or run choice 2 and type them there).')
    say('Press K at the menu to open the keyword list - check the emails you want hidden.')
    return summaries


def build_message_rows(ref, records, pdf_path):
    return [build_row(ref, record, attachments, pdf_path) for record, attachments in records]


def eml_path(ref, index, eml_dir):
    candidate = f'{ref}.eml' if index == 1 else f'{ref}-{index}.eml'
    if not os.path.exists(os.path.join(eml_dir, candidate)):
        return candidate
    counter = 2
    while True:
        candidate = f'{ref}-{index}-{counter}.eml'
        if not os.path.exists(os.path.join(eml_dir, candidate)):
            return candidate
        counter += 1


def archive_units_for(summaries, positions):
    units = []
    used = set()
    for position in positions:
        item = summaries[position - 1]
        thread_no = item.get('thread_no') or 0
        if thread_no:
            if thread_no in used:
                continue
            used.add(thread_no)
            unit = [member for member in summaries if (member.get('thread_no') or 0) == thread_no]
        else:
            unit = [item]
        units.append(unit)
    return units


def archive_unit(connection, config, root, unit, seen):
    existing_refs = [seen[member['key']] for member in unit if member['key'] in seen]
    ref = min(existing_refs) if existing_refs else next_ref_no(root)
    records = []
    new_attachments = []
    for index, item in enumerate(unit, 1):
        raw = fetch_full(connection, item['uid'])
        try:
            message = message_from_bytes(raw, policy=policy.default)
        except Exception as exc:
            raise ArchiveError(f'cannot parse the downloaded email ({exc})')
        record = build_record(message, config, fallback_date=item.get('date', ''))
        if item['key'] in seen:
            attachments = []
        else:
            attachments = save_attachments(message, os.path.join(root, 'Attachments'), ref, record['date'])
            new_attachments.extend(attachments)
        records.append((record, attachments))
        if config.get('save_eml') and item['key'] not in seen:
            eml_name = eml_path(ref, index, os.path.join(root, 'EML'))
            try:
                with open(os.path.join(root, 'EML', eml_name), 'wb') as handle:
                    handle.write(raw)
            except OSError as exc:
                warn(f'could not save the .eml backup for {ref}: {exc}')
    make_pdf(os.path.join(root, 'PDF', ref + '.pdf'), ref, records)
    added = 0
    for item, (record, attachments) in zip(unit, records):
        if item['key'] not in seen:
            append_log_row(root, build_row(ref, record, attachments, f'PDF/{ref}.pdf'))
            added += 1
    for item in unit:
        seen[item['key']] = ref
    save_seen(root, seen)
    return {'ref': ref, 'attachments': new_attachments, 'messages': len(unit), 'added': added}


def archive_selected(config, root, summaries, positions, notify):
    units = archive_units_for(summaries, positions)
    seen = load_seen(root)
    notify(f'Connecting to {config.get("imap_server")}...')
    connection = connect(config)
    archived = skipped = failed = 0
    try:
        select_inbox(connection)
        total = len(units)
        for step, unit in enumerate(units, 1):
            head = unit[0]
            sender = head['from_name'] or head['from_email'] or '(unknown sender)'
            thread_text = f' ({len(unit)} messages)' if len(unit) > 1 else ''
            notify(f'[{step}/{total}] {head["date"] or "-"} | {sender} | {head["subject"] or "(no subject)"}{thread_text}')
            if all(member['key'] in seen for member in unit):
                notify(f'      already archived as {seen[head["key"]]} - skipped')
                skipped += 1
                continue
            try:
                result = archive_unit(connection, config, root, unit, seen)
                archived += 1
                if result['added'] < result['messages']:
                    notify(f'      extended as {result["ref"]} | added {result["added"]} of '
                           f'{result["messages"]} message(s) | {len(result["attachments"])} '
                           f'attachment(s) | PDF updated | {result["added"]} log row(s) added')
                else:
                    notify(f'      archived as {result["ref"]} | {result["messages"]} message(s) | '
                           f'{len(result["attachments"])} attachment(s) | PDF saved | '
                           f'{result["added"]} log row(s) added')
            except ArchiveError as exc:
                failed += 1
                fail(f'      {exc}')
            except Exception as exc:
                failed += 1
                fail(f'      unexpected error: {exc}')
    finally:
        close_conn(connection)
    notify(f'Finished: {archived} archived, {skipped} skipped (already archived), {failed} failed.')
    if archived:
        try:
            export_excel(root)
        except ArchiveError as exc:
            warn(str(exc))
    return archived, skipped, failed


def cmd_archive(config, root, summaries):
    if not summaries:
        say('No email list is loaded yet - fetching from the server first.')
        summaries = cmd_fetch(config, root)
        if not summaries:
            return []
    while True:
        selection = ask('Select emails to archive (e.g. 1,3,5, t2 for a whole thread, or "all"; blank to cancel): ').strip()
        if not selection:
            say('Cancelled - nothing was archived.')
            return summaries
        try:
            threads = thread_positions(summaries) if config.get('thread_view', True) else None
            positions = parse_selection(selection, len(summaries), threads)
            break
        except ValueError as exc:
            warn(str(exc))
    if not positions:
        say('Nothing selected - nothing was archived.')
        return summaries
    archive_selected(config, root, summaries, positions, say)
    return summaries


def run_menu(config, root, config_path):
    say('=' * 66)
    say('  Email Archive Tool')
    say(f'  Account:  {config.get("email")}')
    say(f'  Archive:  {root}')
    say('=' * 66)
    summaries = []
    while True:
        choice = ask('Choose  [1] fetch  [2] archive  [3] Excel  [4] quit  (k)eywords: ')
        choice = choice.strip()
        if choice == '1':
            try:
                summaries = cmd_fetch(config, root)
            except ArchiveError as exc:
                fail(str(exc))
        elif choice == '2':
            try:
                summaries = cmd_archive(config, root, summaries)
            except ArchiveError as exc:
                fail(str(exc))
        elif choice == '3':
            try:
                export_excel(root)
            except ArchiveError as exc:
                fail(str(exc))
        elif choice in ('k', 'K'):
            try:
                cmd_keywords(config_path, config, root, summaries)
            except ArchiveError as exc:
                fail(str(exc))
        elif choice in ('4', 'q', 'quit', 'exit'):
            say('Goodbye.')
            return 0
        else:
            warn('Unknown choice - enter 1, 2, 3, 4 or k.')


def main(argv=None):
    setup_console()
    parser = argparse.ArgumentParser(
        description='Archive selected Outlook.com emails to PDF, attachments, a CSV log and Excel.'
    )
    parser.add_argument('--config', metavar='PATH', default=None,
                        help='path to config.json (default: config.json next to this script)')
    args = parser.parse_args(argv)
    config_path = os.path.abspath(args.config) if args.config else default_config_path()
    try:
        if not check_dependencies():
            return 1
        config = load_config(config_path)
        root = resolve_root(config)
        ensure_dirs(root)
        ensure_credentials(config, config_path)
        return run_menu(config, root, config_path)
    except ArchiveError as exc:
        fail(str(exc))
        return 1
    except KeyboardInterrupt:
        say('')
        warn('Cancelled by the user.')
        return 130


if __name__ == '__main__':
    sys.exit(main())
